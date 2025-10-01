"""
ISA_Resume_Builder Agent - AI-Powered Resume Optimization
Google Vertex AI agent for building tailored resumes based on job requirements and candidate experience.
"""

import asyncio
import logging
from typing import Dict, Any, List, Optional
from datetime import datetime
import json
import re
import base64
import io

from google.cloud import aiplatform
import vertexai
from vertexai.generative_models import GenerativeModel

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available. DOCX export will be disabled.")

import sys
import os
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import settings

logger = logging.getLogger(__name__)


class ISAResumeBuilder:
    """
    ISA_Resume_Builder Agent
    
    Expert recruiter agent that builds optimized resumes by:
    - Analyzing job descriptions for keywords and requirements
    - Retrieving candidate's current resume and interview transcripts
    - Tailoring experience to match specific job requirements
    - Optimizing for ATS (Applicant Tracking Systems)
    - Providing keyword matching and improvement suggestions
    """
    
    def __init__(self):
        """Initialize the ISA_Resume_Builder agent"""
        self.project_id = settings.GOOGLE_CLOUD_PROJECT_ID
        self.location = settings.VERTEX_AI_LOCATION
        self.model_name = settings.VERTEX_AI_MODEL
        
        # Initialize Vertex AI
        if self.project_id:
            vertexai.init(project=self.project_id, location=self.location)
        
        self.model = None
        self._initialize_model()
    
    def _initialize_model(self):
        """Initialize the Vertex AI model"""
        try:
            if self.project_id:
                self.model = GenerativeModel(self.model_name)
                logger.info(f"✅ ISA_Resume_Builder model initialized: {self.model_name}")
            else:
                logger.warning("⚠️ Google Cloud Project ID not set - ISA_Resume_Builder will run in mock mode")
        except Exception as e:
            logger.error(f"❌ Failed to initialize ISA_Resume_Builder model: {e}")
            self.model = None
    
    async def build_resume(
        self,
        job_description: str,
        company_name: str,
        job_title: str,
        candidate_resume: str,
        opportunity_transcripts: List[str] = None,
        interview_transcripts: List[str] = None,
        contact_name: str = "",
        user_id: str = ""
    ) -> Dict[str, Any]:
        """
        Build an optimized resume tailored to the specific job
        
        Args:
            job_description: Full job description and requirements
            company_name: Target company name
            job_title: Target job title
            candidate_resume: Current candidate resume text
            opportunity_transcripts: List of opportunity discussion transcripts
            interview_transcripts: List of THEIA interview transcripts
            contact_name: Candidate's name from Salesforce
            user_id: User/Contact ID for logging
            
        Returns:
            Dict containing optimized resume and analysis
        """
        try:
            logger.info(f"🚀 ISA_Resume_Builder starting for user {user_id}: {company_name} - {job_title}")
            
            if not self.model:
                return self._mock_resume_response(job_description, company_name, job_title)
            
            # Prepare all available candidate data
            candidate_data = self._prepare_candidate_data(
                candidate_resume, 
                opportunity_transcripts or [], 
                interview_transcripts or [],
                contact_name
            )
            
            # Analyze job requirements
            job_analysis = await self._analyze_job_requirements(job_description, company_name, job_title)
            
            # Build optimized resume
            optimized_resume = await self._build_optimized_resume(
                job_analysis, 
                candidate_data, 
                company_name, 
                job_title
            )
            
            # Calculate matching metrics
            metrics = await self._calculate_resume_metrics(optimized_resume, job_analysis)
            
            # Generate DOCX if available
            resume_docx_base64 = None
            if DOCX_AVAILABLE:
                try:
                    resume_docx_base64 = self._generate_resume_docx(optimized_resume, contact_name or "Resume")
                except Exception as e:
                    logger.warning(f"DOCX generation failed: {e}")

            result = {
                "status": "success",
                "resume_text": optimized_resume,
                "resume_html": self._format_resume_html(optimized_resume),
                "resume_docx": resume_docx_base64,
                "company_name": company_name,
                "job_title": job_title,
                "keywords_matched": metrics.get("keywords_matched", 0),
                "ats_score": metrics.get("ats_score", "N/A"),
                "match_percentage": metrics.get("match_percentage", 0),
                "optimization_notes": metrics.get("optimization_notes", []),
                "job_keywords": job_analysis.get("keywords", []),
                "timestamp": datetime.utcnow().isoformat()
            }
            
            logger.info(f"✅ ISA_Resume_Builder completed for user {user_id}: {metrics.get('match_percentage', 0)}% match")
            return result
            
        except Exception as e:
            logger.error(f"❌ ISA_Resume_Builder failed for user {user_id}: {e}")
            return {
                "status": "error",
                "error": str(e),
                "resume_text": candidate_resume or "",
                "company_name": company_name,
                "job_title": job_title,
                "timestamp": datetime.utcnow().isoformat()
            }
    
    def _prepare_candidate_data(
        self, 
        resume: str, 
        opportunity_transcripts: List[str], 
        interview_transcripts: List[str],
        contact_name: str
    ) -> Dict[str, Any]:
        """Prepare and structure candidate data from all sources"""
        
        # Extract experience from transcripts
        experience_from_transcripts = []
        
        for transcript in opportunity_transcripts:
            if transcript and transcript.strip():
                experience_from_transcripts.append({
                    "source": "opportunity_discussion",
                    "content": transcript[:2000]  # Limit length
                })
        
        for transcript in interview_transcripts:
            if transcript and transcript.strip():
                experience_from_transcripts.append({
                    "source": "theia_interview",
                    "content": transcript[:2000]  # Limit length
                })
        
        return {
            "contact_name": contact_name or "Candidate",
            "current_resume": resume or "",
            "additional_experience": experience_from_transcripts,
            "total_sources": 1 + len(experience_from_transcripts)
        }
    
    async def _analyze_job_requirements(self, job_description: str, company_name: str, job_title: str) -> Dict[str, Any]:
        """Analyze job description to extract key requirements and keywords"""
        
        prompt = f"""
        As an expert recruiter and ATS specialist, analyze this job description for a {job_title} position at {company_name}.

        Job Description:
        {job_description}

        Extract and provide:
        1. Key required skills and technologies (prioritized list)
        2. Important keywords for ATS optimization
        3. Required experience level and years
        4. Educational requirements
        5. Soft skills mentioned
        6. Industry-specific terms
        7. Company culture keywords
        8. Action verbs commonly used in the posting

        Format your response as JSON with the following structure:
        {{
            "required_skills": ["skill1", "skill2", ...],
            "keywords": ["keyword1", "keyword2", ...],
            "experience_level": "entry/mid/senior",
            "years_experience": "X-Y years",
            "education": ["degree1", "degree2", ...],
            "soft_skills": ["skill1", "skill2", ...],
            "industry_terms": ["term1", "term2", ...],
            "company_culture": ["value1", "value2", ...],
            "action_verbs": ["verb1", "verb2", ...],
            "priority_keywords": ["top10", "keywords", "for", "ats"]
        }}
        """
        
        try:
            response = await asyncio.get_event_loop().run_in_executor(
                None, 
                lambda: self.model.generate_content(prompt)
            )
            
            # Extract JSON from response
            response_text = response.text
            json_match = re.search(r'\{.*\}', response_text, re.DOTALL)
            
            if json_match:
                return json.loads(json_match.group())
            else:
                logger.warning("Failed to parse job analysis JSON, using fallback")
                return self._fallback_job_analysis(job_description)
                
        except Exception as e:
            logger.error(f"Job analysis failed: {e}")
            return self._fallback_job_analysis(job_description)
    
    def _fallback_job_analysis(self, job_description: str) -> Dict[str, Any]:
        """Fallback job analysis using basic text processing"""
        
        # Simple keyword extraction
        common_skills = [
            "Python", "Java", "JavaScript", "SQL", "AWS", "Docker", "Kubernetes",
            "React", "Node.js", "Git", "Agile", "Scrum", "REST API", "Machine Learning",
            "Data Analysis", "Project Management", "Leadership", "Communication"
        ]
        
        found_skills = [skill for skill in common_skills if skill.lower() in job_description.lower()]
        
        return {
            "required_skills": found_skills[:10],
            "keywords": found_skills[:15],
            "experience_level": "mid",
            "years_experience": "3-5 years",
            "education": ["Bachelor's degree"],
            "soft_skills": ["Communication", "Teamwork", "Problem-solving"],
            "industry_terms": [],
            "company_culture": [],
            "action_verbs": ["developed", "implemented", "managed", "led"],
            "priority_keywords": found_skills[:10]
        }
    
    async def _build_optimized_resume(
        self, 
        job_analysis: Dict[str, Any], 
        candidate_data: Dict[str, Any], 
        company_name: str, 
        job_title: str
    ) -> str:
        """Build optimized resume using job analysis and candidate data"""
        
        prompt = f"""
        As an expert recruiter specializing in ATS optimization, create a tailored resume for a {job_title} position at {company_name}.

        TARGET JOB REQUIREMENTS:
        - Required Skills: {', '.join(job_analysis.get('required_skills', []))}
        - Priority Keywords: {', '.join(job_analysis.get('priority_keywords', []))}
        - Experience Level: {job_analysis.get('experience_level', 'mid')}
        - Education: {', '.join(job_analysis.get('education', []))}

        CANDIDATE DATA:
        Current Resume:
        {candidate_data.get('current_resume', 'No resume provided')}

        Additional Experience from Interviews/Discussions:
        {chr(10).join([f"- {exp['content'][:500]}..." for exp in candidate_data.get('additional_experience', [])])}

        CRITICAL FACTUAL REQUIREMENTS:
        1. ONLY use skills, experience, and achievements that are explicitly mentioned in the candidate's resume or interview transcripts
        2. DO NOT add or invent any skills, certifications, or experiences not present in the source material
        3. DO NOT include industry buzzwords unless they appear in the candidate's actual background
        4. If a job requirement doesn't match the candidate's experience, DO NOT fabricate it
        5. Use exact company names, job titles, and dates from the candidate's resume
        6. Only mention technologies, tools, or methodologies the candidate has actually used

        FORMATTING REQUIREMENTS:
        - NO introductory text or explanations
        - NO markdown symbols (###, **, ---)
        - Use clean, professional formatting with single spacing
        - Start directly with the candidate's name in BOLD and CENTERED
        - Use CAPITAL LETTERS for section headers
        - Use bullet points (•) for lists with single spacing between items
        - Keep skill lists concise (max 8-10 items per section)
        - Limit resume to 1-2 pages maximum

        RESUME STRUCTURE:
        - Candidate Name (BOLD, CENTERED)
        - Contact Information (centered)
        - Professional Summary (2-3 concise lines)
        - Core Competencies (max 10 skills from actual experience)
        - Professional Experience (reverse chronological, actual positions only)
        - Education (actual degrees/certifications only)

        CONTENT OPTIMIZATION:
        - Emphasize candidate's actual strengths that align with job requirements
        - Use action verbs to describe real accomplishments
        - Quantify achievements where data is available in source material
        - Ensure ATS optimization through natural keyword placement of verified skills

        Generate ONLY the resume content, starting directly with the candidate's name in BOLD and CENTERED:
        """
        
        try:
            response = await asyncio.get_event_loop().run_in_executor(
                None, 
                lambda: self.model.generate_content(prompt)
            )
            
            raw_text = response.text.strip()
            cleaned_resume = self._clean_resume_text(raw_text)
            return cleaned_resume
            
        except Exception as e:
            logger.error(f"Resume generation failed: {e}")
            return self._fallback_resume(candidate_data, job_analysis, company_name, job_title)
    
    def _fallback_resume(
        self, 
        candidate_data: Dict[str, Any], 
        job_analysis: Dict[str, Any], 
        company_name: str, 
        job_title: str
    ) -> str:
        """Generate factual fallback resume when AI generation fails"""
        
        name = candidate_data.get('contact_name', 'CANDIDATE NAME')
        current_resume = candidate_data.get('current_resume', '')
        
        # Only use skills that appear in the candidate's actual resume
        actual_skills = []
        if current_resume:
            # Extract skills that are actually mentioned in the resume
            resume_lower = current_resume.lower()
            for skill in job_analysis.get('required_skills', []):
                if skill.lower() in resume_lower:
                    actual_skills.append(skill)
        
        # If no matching skills found, provide minimal factual resume
        if not actual_skills:
            return f"""
{name.upper()}

PROFESSIONAL SUMMARY
Professional with relevant experience seeking opportunities to contribute to {company_name}.

Note: Please provide your detailed resume and experience information to generate a comprehensive, tailored resume for the {job_title} position.

EXPERIENCE
As detailed in your provided resume and interview responses.

EDUCATION
As specified in your background information.
        """.strip()
        
        # Generate resume with only verified information
        skills_text = ' • '.join(actual_skills[:8])  # Limit to 8 skills max
        
        return f"""
{name.upper()}

PROFESSIONAL SUMMARY
Professional with experience relevant to {job_title} opportunities, with demonstrated expertise in areas that align with {company_name}'s requirements.

CORE COMPETENCIES
• {skills_text}

PROFESSIONAL EXPERIENCE
[Based on information provided in resume and interview transcripts]

EDUCATION
[As specified in candidate background]

Note: This is a condensed version. Complete details available in provided resume and interview responses.
        """.strip()
    
    async def _calculate_resume_metrics(self, resume: str, job_analysis: Dict[str, Any]) -> Dict[str, Any]:
        """Calculate resume matching metrics and provide optimization suggestions"""
        
        priority_keywords = job_analysis.get('priority_keywords', [])
        all_keywords = job_analysis.get('keywords', [])
        
        # Count keyword matches (case-insensitive)
        resume_lower = resume.lower()
        priority_matches = sum(1 for keyword in priority_keywords if keyword.lower() in resume_lower)
        total_matches = sum(1 for keyword in all_keywords if keyword.lower() in resume_lower)
        
        # Calculate match percentage
        match_percentage = int((priority_matches / max(len(priority_keywords), 1)) * 100)
        
        # Generate ATS score
        ats_score = min(95, 60 + (match_percentage * 0.3) + (total_matches * 2))
        ats_score = f"{int(ats_score)}/100"
        
        # Generate optimization notes
        missing_keywords = [kw for kw in priority_keywords if kw.lower() not in resume_lower]
        optimization_notes = []
        
        if missing_keywords:
            optimization_notes.append(f"Consider adding these key terms: {', '.join(missing_keywords[:5])}")
        
        if match_percentage < 70:
            optimization_notes.append("Resume could benefit from more job-specific keywords")
        
        if len(resume.split()) < 200:
            optimization_notes.append("Resume might be too brief - consider adding more detail")
        
        if not any(verb in resume_lower for verb in ["developed", "implemented", "managed", "led", "created"]):
            optimization_notes.append("Add more action verbs to strengthen impact statements")
        
        return {
            "keywords_matched": total_matches,
            "priority_keywords_matched": priority_matches,
            "ats_score": ats_score,
            "match_percentage": match_percentage,
            "optimization_notes": optimization_notes
        }
    
    def _format_resume_html(self, resume_text: str) -> str:
        """Convert plain text resume to rich HTML format with optimized spacing"""
        
        if not resume_text:
            return "<p>No resume content available</p>"
        
        # Enhanced HTML formatting with optimized spacing
        lines = resume_text.split('\n')
        html_lines = []
        
        # Add container div with professional styling and tighter line spacing
        html_lines.append('<div style="font-family: \'Times New Roman\', serif; line-height: 1.3; color: #333; max-width: 800px;">')
        
        current_section = None
        
        for line in lines:
            line = line.strip()
            if not line:
                # Reduced empty line spacing
                html_lines.append('<div style="margin: 3px 0;"></div>')
                continue
            
            # Detect name/header - ensure it's bold and centered
            if self._is_name_line(line):
                html_lines.append(f'<h1 style="text-align: center; font-size: 24px; font-weight: bold; margin: 0 0 8px 0; color: #1a202c; letter-spacing: 1px; text-transform: uppercase;">{line}</h1>')
            
            # Detect contact info (contains email, phone, etc.)
            elif self._is_contact_line(line):
                html_lines.append(f'<p style="text-align: center; margin: 3px 0 15px 0; color: #4a5568; font-size: 14px;">{line}</p>')
            
            # Section headers (ALL CAPS) - reduced spacing
            elif line.isupper() and len(line.split()) <= 4:
                current_section = line
                html_lines.append(f'<h2 style="border-bottom: 1px solid #2d3748; padding-bottom: 3px; margin: 15px 0 8px 0; font-size: 16px; font-weight: bold; color: #2d3748; text-transform: uppercase; letter-spacing: 0.5px;">{line}</h2>')
            
            # Bullet points - single spacing
            elif line.startswith('•') or line.startswith('-') or line.startswith('*'):
                bullet_content = line[1:].strip() if line[0] in '•-*' else line
                html_lines.append(f'<div style="margin: 2px 0 2px 20px; position: relative; line-height: 1.2;"><span style="position: absolute; left: -15px; color: #2d3748;">•</span>{bullet_content}</div>')
            
            # Job titles/company names - reduced spacing
            elif '|' in line or self._contains_dates(line):
                if '|' in line:
                    parts = line.split('|')
                    if len(parts) >= 2:
                        job_title = parts[0].strip()
                        company_info = ' | '.join(parts[1:]).strip()
                        html_lines.append(f'<div style="margin: 10px 0 3px 0;"><strong style="color: #2d3748; font-size: 15px;">{job_title}</strong></div>')
                        html_lines.append(f'<div style="color: #4a5568; font-style: italic; margin-bottom: 5px; font-size: 13px;">{company_info}</div>')
                    else:
                        html_lines.append(f'<div style="margin: 10px 0 3px 0; font-weight: bold; color: #2d3748;">{line}</div>')
                else:
                    html_lines.append(f'<div style="margin: 10px 0 3px 0; font-weight: bold; color: #2d3748;">{line}</div>')
            
            # Regular content - reduced spacing
            else:
                html_lines.append(f'<p style="margin: 4px 0; text-align: justify; line-height: 1.2;">{line}</p>')
        
        html_lines.append('</div>')
        return '\n'.join(html_lines)
    
    def _is_name_line(self, line: str) -> bool:
        """Check if line is likely the candidate's name"""
        # Usually the first non-empty line or a line in ALL CAPS that looks like a name
        words = line.split()
        return (len(words) >= 2 and len(words) <= 4 and 
                line.isupper() and 
                not any(char in line for char in ['@', '|', '•', '-', '(', ')']))
    
    def _is_contact_line(self, line: str) -> bool:
        """Check if line contains contact information"""
        contact_indicators = ['@', 'phone:', 'email:', 'linkedin', '(', ')', '-', '•']
        return any(indicator in line.lower() for indicator in contact_indicators)
    
    def _contains_dates(self, line: str) -> bool:
        """Check if line contains date patterns"""
        import re
        date_patterns = [
            r'\d{4}\s*-\s*\d{4}',  # 2020 - 2024
            r'\d{4}\s*-\s*Present',  # 2020 - Present
            r'\w+\s+\d{4}',  # Jan 2020
        ]
        return any(re.search(pattern, line) for pattern in date_patterns)
    
    def _generate_resume_docx(self, resume_text: str, filename_base: str = "Resume") -> str:
        """Generate a DOCX version of the resume and return as base64 string"""
        
        if not DOCX_AVAILABLE:
            logger.warning("DOCX generation requested but python-docx not available")
            return None
        
        try:
            # Create new document
            doc = Document()
            
            # Set document margins
            sections = doc.sections
            for section in sections:
                section.top_margin = Inches(0.5)
                section.bottom_margin = Inches(0.5)
                section.left_margin = Inches(0.75)
                section.right_margin = Inches(0.75)
            
            lines = resume_text.split('\n')
            
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                # Detect name/header - ensure bold and centered
                if self._is_name_line(line):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.space_before = Pt(0)
                    p.space_after = Pt(6)
                    run = p.add_run(line.upper())  # Ensure name is uppercase
                    run.font.size = Pt(18)
                    run.bold = True
                    run.font.name = 'Times New Roman'
                
                # Detect contact info
                elif self._is_contact_line(line):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(line)
                    run.font.size = Pt(11)
                    run.font.name = 'Times New Roman'
                
                # Section headers (ALL CAPS) - reduced spacing
                elif line.isupper() and len(line.split()) <= 4:
                    p = doc.add_paragraph()
                    p.space_before = Pt(8)
                    p.space_after = Pt(3)
                    run = p.add_run(line)
                    run.font.size = Pt(12)
                    run.bold = True
                    run.font.name = 'Times New Roman'
                    run.underline = True  # Simple underline instead of separate paragraph
                
                # Bullet points - single spacing
                elif line.startswith('•') or line.startswith('-') or line.startswith('*'):
                    p = doc.add_paragraph()
                    p.style = 'List Bullet'
                    p.space_before = Pt(1)
                    p.space_after = Pt(1)
                    bullet_content = line[1:].strip() if line[0] in '•-*' else line
                    run = p.add_run(bullet_content)
                    run.font.size = Pt(11)
                    run.font.name = 'Times New Roman'
                
                # Job titles/company info - reduced spacing
                elif '|' in line or self._contains_dates(line):
                    p = doc.add_paragraph()
                    p.space_before = Pt(4)
                    if '|' in line:
                        parts = line.split('|')
                        # Job title in bold
                        run = p.add_run(parts[0].strip())
                        run.bold = True
                        run.font.size = Pt(11)
                        run.font.name = 'Times New Roman'
                        
                        if len(parts) > 1:
                            # Company info in regular text
                            run = p.add_run(' | ' + ' | '.join(parts[1:]).strip())
                            run.font.size = Pt(11)
                            run.font.name = 'Times New Roman'
                    else:
                        run = p.add_run(line)
                        run.bold = True
                        run.font.size = Pt(11)
                        run.font.name = 'Times New Roman'
                
                # Regular text
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(line)
                    run.font.size = Pt(11)
                    run.font.name = 'Times New Roman'
            
            # Save to memory buffer
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            # Convert to base64
            docx_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            logger.info(f"Successfully generated DOCX resume ({len(docx_base64)} chars base64)")
            return docx_base64
            
        except Exception as e:
            logger.error(f"Failed to generate DOCX: {e}")
            return None
    
    def _clean_resume_text(self, resume_text: str) -> str:
        """Clean and format resume text by removing unwanted elements"""
        
        if not resume_text:
            return ""
        
        lines = resume_text.split('\n')
        cleaned_lines = []
        
        # Skip common intro phrases
        skip_patterns = [
            "of course",
            "here is a tailored resume",
            "here's a tailored resume", 
            "below is a tailored resume",
            "i've created a tailored resume",
            "here is the optimized resume",
            "here's the optimized resume"
        ]
        
        start_processing = False
        
        for line in lines:
            line_lower = line.lower().strip()
            
            # Skip intro lines
            if not start_processing:
                # Check if this line should be skipped
                skip_line = any(pattern in line_lower for pattern in skip_patterns)
                if skip_line or line_lower in ["---", "```", ""]:
                    continue
                # Once we find actual content, start processing
                if line.strip():
                    start_processing = True
            
            if start_processing:
                # Clean the line
                cleaned_line = self._clean_line(line)
                if cleaned_line.strip():  # Only add non-empty lines
                    cleaned_lines.append(cleaned_line)
        
        return '\n'.join(cleaned_lines).strip()
    
    def _clean_line(self, line: str) -> str:
        """Clean individual line by removing markdown symbols"""
        
        # Remove markdown headers (### text)
        line = re.sub(r'^#+\s*', '', line)
        
        # Remove bold markdown (**text**)
        line = re.sub(r'\*\*(.*?)\*\*', r'\1', line)
        
        # Remove italic markdown (*text*)
        line = re.sub(r'\*(.*?)\*', r'\1', line)
        
        # Remove horizontal rules
        if line.strip() in ['---', '***', '___']:
            return ''
        
        # Clean up multiple spaces
        line = re.sub(r'\s+', ' ', line)
        
        return line.strip()
    
    def _mock_resume_response(self, job_description: str, company_name: str, job_title: str) -> Dict[str, Any]:
        """Mock response when Vertex AI is not available"""
        
        mock_resume = f"""
JOHN DOE
Email: john.doe@email.com | Phone: (555) 123-4567 | LinkedIn: linkedin.com/in/johndoe

PROFESSIONAL SUMMARY
Experienced {job_title} with strong background in relevant technologies and methodologies. 
Proven track record of delivering high-quality solutions and driving business results.

CORE COMPETENCIES
• Software Development  • Project Management  • Team Leadership
• Problem Solving      • Communication      • Technical Analysis

PROFESSIONAL EXPERIENCE

Senior {job_title} | Tech Company | 2020 - Present
• Led development of key features and improvements
• Collaborated with cross-functional teams to deliver projects
• Mentored junior team members and improved processes

{job_title} | Previous Company | 2018 - 2020
• Developed and maintained software applications
• Participated in agile development processes
• Contributed to technical documentation and best practices

EDUCATION
Bachelor of Science in Computer Science
State University | 2018

ADDITIONAL SKILLS
Relevant to {company_name} requirements and {job_title} position
        """
        
        # Generate DOCX for mock response if available
        resume_docx_base64 = None
        if DOCX_AVAILABLE:
            try:
                resume_docx_base64 = self._generate_resume_docx(mock_resume.strip(), "Sample Resume")
            except Exception as e:
                logger.warning(f"Mock DOCX generation failed: {e}")

        return {
            "status": "success",
            "resume_text": mock_resume.strip(),
            "resume_html": self._format_resume_html(mock_resume.strip()),
            "resume_docx": resume_docx_base64,
            "company_name": company_name,
            "job_title": job_title,
            "keywords_matched": 8,
            "ats_score": "75/100",
            "match_percentage": 75,
            "optimization_notes": [
                "This is a mock resume generated without AI",
                "Configure Vertex AI for full functionality",
                "Add specific keywords from the job description"
            ],
            "job_keywords": ["software", "development", "management", "leadership"],
            "timestamp": datetime.utcnow().isoformat()
        }


# Export for use in other modules
__all__ = ['ISAResumeBuilder']
